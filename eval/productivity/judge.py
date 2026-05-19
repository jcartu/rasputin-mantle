from __future__ import annotations

import asyncio
import json
import re
import sys
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
GATEWAY_SRC = ROOT / "apps" / "gateway" / "src"
if str(GATEWAY_SRC) not in sys.path:
    sys.path.insert(0, str(GATEWAY_SRC))

from gateway.model_client import anthropic_chat  # noqa: E402

MODEL_NAME = "claude-sonnet-4-6"
RUBRICS_DIR = Path(__file__).with_name("rubrics")
MAX_EXTRACTED_CHARS = 24_000


def judge_structure(task: dict[str, Any], produced_files: list[Path], metadata: dict[str, Any]) -> dict[str, Any]:
    extracted = extract_file_content(produced_files)
    structural = _validate_structure(task, produced_files, metadata, extracted)
    quality = _judge_quality(task, extracted)
    score = min(1.0, structural["score"] + quality["score"])
    return {
        "judge_model": MODEL_NAME,
        "file_match": structural["file_match"],
        "structural_match": structural["structural_match"],
        "structural_score": structural["score"],
        "quality_score_10": quality["quality_score_10"],
        "quality_score": quality["score"],
        "score": score,
        "passed_threshold": 0.70,
        "rubric": quality["rubric"],
        "anthropic": quality["anthropic"],
        "structural_checks": structural["checks"],
        "extracted_content_chars": len(extracted),
    }


def extract_file_content(produced_files: list[Path]) -> str:
    parts: list[str] = []
    for path in produced_files:
        if path.suffix == ".pptx":
            parts.append(_extract_pptx(path))
        elif path.suffix == ".xlsx":
            parts.append(_extract_xlsx(path))
        elif path.suffix == ".docx":
            parts.append(_extract_docx(path))
        elif path.suffix == ".pdf":
            parts.append(_extract_pdf(path))
        else:
            parts.append(f"# {path.name}\nUnsupported file type for text extraction: {path.suffix}")
    return "\n\n".join(parts)[:MAX_EXTRACTED_CHARS]


def _validate_structure(
    task: dict[str, Any], produced_files: list[Path], metadata: dict[str, Any], extracted: str
) -> dict[str, Any]:
    expected = dict(task.get("expected") or {})
    names = {path.name for path in produced_files}
    required = set(expected.get("files") or ([expected["file"]] if expected.get("file") else []))
    file_match = required.issubset(names)
    checks: dict[str, bool] = {
        "file_exists": file_match,
        "opens_cleanly": bool(metadata.get("opens_cleanly", False)),
    }

    if expected.get("min_slides") is not None:
        checks["min_slides"] = metadata.get("slide_count", 0) >= int(expected["min_slides"])
    if expected.get("min_sheets") is not None:
        checks["min_sheets"] = len(metadata.get("sheets", [])) >= int(expected["min_sheets"])
    if expected.get("sheets") is not None:
        checks["required_sheets"] = set(expected["sheets"]).issubset(set(metadata.get("sheets", [])))
    if expected.get("formulas"):
        checks["has_formulas"] = bool(metadata.get("has_formulas")) or _workbooks_have_formulas(produced_files)
    if expected.get("sections"):
        checks["required_sections"] = _content_has_all(extracted, [str(section) for section in expected["sections"]])
    if expected.get("required_headers"):
        checks["required_headers"] = _workbooks_have_headers(
            produced_files, [str(header) for header in expected["required_headers"]]
        )
    if expected.get("max_rows") is not None:
        checks["max_rows"] = _workbooks_within_row_limit(produced_files, int(expected["max_rows"]))
    if expected.get("toc"):
        checks["toc"] = "table of contents" in extracted.lower()
    if expected.get("charts") or _task_requests_charts(task):
        checks["has_charts"] = _files_have_charts(produced_files)

    structural_match = all(checks.values())
    return {
        "file_match": file_match,
        "structural_match": structural_match,
        "score": 0.5 if structural_match else 0.0,
        "checks": checks,
    }


def _judge_quality(task: dict[str, Any], extracted: str) -> dict[str, Any]:
    rubric_name = str(task["format"])
    rubric = _load_rubric(rubric_name)
    prompt = {
        "instruction": (
            "You are judging a productivity artifact. Return JSON only with keys "
            "quality_score and reason. quality_score must be a number from 0 to 10. "
            "Judge against the requested scope, not an imagined enterprise deliverable. "
            "For deliberately small or minimal prompts, do not require invented facts, dates, "
            "or metrics that were not requested. A thin but complete artifact should be in "
            "the 4-6 range; reserve 0-3 for missing, incoherent, or unusable artifacts."
        ),
        "task_prompt": task.get("prompt", ""),
        "task_input": task.get("input", {}),
        "rubric": rubric,
        "extracted_file_content": extracted or "No extractable content was produced.",
    }
    response = asyncio.run(
        anthropic_chat(
            model=MODEL_NAME,
            messages=[{"role": "user", "content": json.dumps(prompt, ensure_ascii=False)}],
            max_tokens=512,
            temperature=0,
            workspace_id="productivity-bench-judge",
        )
    )
    quality_score_10 = _parse_quality_score(response["content"])
    return {
        "score": quality_score_10 / 10.0 * 0.5,
        "quality_score_10": quality_score_10,
        "rubric": rubric_name,
        "anthropic": {
            "model": response["model"],
            "input_tokens": response["input_tokens"],
            "output_tokens": response["output_tokens"],
            "cost_usd": response["cost_usd"],
            "latency_ms": response["latency_ms"],
            "raw_response": response["content"],
        },
    }


def _load_rubric(format_name: str) -> str:
    path = RUBRICS_DIR / f"{format_name}.md"
    if not path.exists():
        raise FileNotFoundError(f"Missing productivity rubric: {path}")
    return path.read_text(encoding="utf-8")


def _parse_quality_score(content: str) -> float:
    try:
        data = json.loads(content)
        value = data.get("quality_score")
    except json.JSONDecodeError:
        match = re.search(r"(?:quality_score|score)\D+(\d+(?:\.\d+)?)", content, flags=re.IGNORECASE)
        if not match:
            match = re.search(r"\b(10(?:\.0+)?|[0-9](?:\.\d+)?)\b", content)
        if not match:
            raise ValueError(f"Could not parse quality score from Sonnet response: {content!r}")
        value = match.group(1)
    score = float(value)
    return max(0.0, min(10.0, score))


def _extract_pptx(path: Path) -> str:
    prs = Presentation(path)
    lines = [f"# {path.name}", f"Slide count: {len(prs.slides)}"]
    for index, slide in enumerate(prs.slides, start=1):
        lines.append(f"## Slide {index}")
        image_count = sum(1 for shape in slide.shapes if shape.shape_type == 13)
        if image_count:
            lines.append(f"Visuals: {image_count} image/chart artifact(s) present")
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = "\n".join(
                paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()
            )
            if text:
                label = "Notes" if shape.name == "Speaker Notes" or text.lower().startswith("notes:") else "Text"
                lines.append(f"{label}: {text}")
    return "\n".join(lines)


def _extract_xlsx(path: Path) -> str:
    wb = load_workbook(path, data_only=False)
    lines = [f"# {path.name}", f"Sheets: {', '.join(wb.sheetnames)}"]
    for sheet in wb.worksheets:
        lines.append(f"## Sheet: {sheet.title}")
        charts = len(getattr(sheet, "_charts", []))
        images = len(getattr(sheet, "_images", []))
        if charts or images:
            lines.append(f"Visuals: {charts} chart object(s), {images} image/chart artifact(s)")
        for row in sheet.iter_rows(max_row=min(sheet.max_row, 30), max_col=min(sheet.max_column, 12)):
            values = [_cell_text(cell.value) for cell in row]
            if any(values):
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _extract_docx(path: Path) -> str:
    document = Document(path)
    lines = [f"# {path.name}"]
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style is not None else ""
        prefix = "Heading" if style.startswith("Heading") or style == "Title" else "Paragraph"
        lines.append(f"{prefix}: {text}")
    for index, table in enumerate(document.tables, start=1):
        lines.append(f"Table {index}:")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        return f"# {path.name}\n{text}" if text else f"# {path.name}\nPDF contained no extractable text."
    except Exception as exc:  # pragma: no cover - optional parser diagnostics
        return f"# {path.name}\nPDF text extraction failed: {exc}"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _content_has_all(content: str, needles: list[str]) -> bool:
    haystack = content.casefold()
    return all(needle.casefold() in haystack for needle in needles)


def _workbooks_have_formulas(paths: list[Path]) -> bool:
    for path in paths:
        if path.suffix != ".xlsx":
            continue
        workbook = load_workbook(path, data_only=False)
        if any(
            isinstance(cell.value, str) and cell.value.startswith("=")
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ):
            return True
    return False


def _workbooks_have_headers(paths: list[Path], headers: list[str]) -> bool:
    expected = {header.casefold() for header in headers}
    found: set[str] = set()
    for path in paths:
        if path.suffix != ".xlsx":
            continue
        workbook = load_workbook(path, data_only=False)
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(max_row=1, values_only=True):
                found.update(str(value).casefold() for value in row if value is not None)
    return expected.issubset(found)


def _workbooks_within_row_limit(paths: list[Path], max_rows: int) -> bool:
    checked = False
    for path in paths:
        if path.suffix != ".xlsx":
            continue
        checked = True
        workbook = load_workbook(path, data_only=False)
        if any(sheet.max_row > max_rows for sheet in workbook.worksheets if sheet.max_row > 1):
            return False
    return checked


def _task_requests_charts(task: dict[str, Any]) -> bool:
    task_input = dict(task.get("input") or {})
    if task_input.get("charts"):
        return True
    outline = dict(task_input.get("outline") or {})
    return any(bool(section.get("chart")) for section in outline.get("sections") or [])


def _files_have_charts(paths: list[Path]) -> bool:
    for path in paths:
        if path.suffix == ".xlsx" and _workbook_has_chart_artifacts(path):
            return True
        if path.suffix == ".pptx" and _presentation_has_images(path):
            return True
    return False


def _workbook_has_chart_artifacts(path: Path) -> bool:
    workbook = load_workbook(path, data_only=False)
    return any(getattr(sheet, "_charts", []) or getattr(sheet, "_images", []) for sheet in workbook.worksheets)


def _presentation_has_images(path: Path) -> bool:
    prs = Presentation(path)
    return any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes)
