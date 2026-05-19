from __future__ import annotations

import argparse
import json
import sys
import tempfile
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT / "packages" / "codeact") not in sys.path:
    sys.path.insert(0, str(ROOT / "packages" / "codeact"))
if str(ROOT / "packages" / "sandbox") not in sys.path:
    sys.path.insert(0, str(ROOT / "packages" / "sandbox"))
if str(Path(__file__).resolve().parent) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parent))

from codeact.skills_loader import invoke_skill  # noqa: E402
from judge import judge_structure  # noqa: E402


def load_tasks(path: Path) -> list[dict[str, Any]]:
    payload = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return list(payload.get("tasks") or [])


def run_task(task: dict[str, Any], root: Path) -> dict[str, Any]:
    skill = {"slides": "slides", "spreadsheet": "spreadsheet", "document": "document"}[str(task["format"])]
    with tempfile.TemporaryDirectory(prefix="mantle-productivity-") as tmp:
        output_dir = Path(tmp)
        args: dict[str, Any] = {}
        args["prompt"] = str(task.get("prompt") or "")
        args["output_dir"] = str(output_dir)
        result = invoke_skill(skill, args)
        produced = sorted(path for path in output_dir.iterdir() if path.is_file())
        metadata = _inspect_outputs(produced)
        judged = judge_structure(task, produced, metadata)
        passed = judged["score"] >= 0.70
        return {
            "id": task["id"],
            "format": task["format"],
            "exit_code": result.exit_code,
            "files": [path.name for path in produced],
            "file_produced": bool(produced),
            "opens_cleanly": metadata.get("opens_cleanly", False),
            "quality": judged,
            "passed": passed,
            "estimated_cost_usd": result.estimated_cost_usd,
            "anthropic_judge_cost_usd": judged.get("anthropic", {}).get("cost_usd", 0.0),
            "stderr": result.stderr,
        }


def _inspect_outputs(paths: list[Path]) -> dict[str, Any]:
    metadata: dict[str, Any] = {"opens_cleanly": True, "sheets": []}
    for path in paths:
        try:
            if path.suffix == ".pptx":
                metadata["slide_count"] = len(Presentation(path).slides)
            elif path.suffix == ".xlsx":
                wb = load_workbook(path, data_only=False)
                metadata["sheets"] = wb.sheetnames
                metadata["has_formulas"] = any(
                    isinstance(cell.value, str) and cell.value.startswith("=")
                    for sheet in wb.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
            elif path.suffix == ".docx":
                document = Document(path)
                metadata["docx_paragraphs"] = len(document.paragraphs)
            elif path.suffix == ".pdf":
                metadata["pdf_bytes"] = path.stat().st_size
        except Exception as exc:  # pragma: no cover - diagnostic payload for benchmark users
            metadata["opens_cleanly"] = False
            metadata["open_error"] = str(exc)
    return metadata


def summarize(results: list[dict[str, Any]]) -> dict[str, Any]:
    by_format: dict[str, list[dict[str, Any]]] = {}
    for result in results:
        by_format.setdefault(str(result["format"]), []).append(result)
    format_rates = {
        name: sum(1 for item in items if item["passed"]) / len(items)
        for name, items in sorted(by_format.items())
        if items
    }
    aggregate = sum(1 for result in results if result["passed"]) / len(results) if results else 0.0
    return {
        "aggregate_pass_rate": aggregate,
        "format_pass_rates": format_rates,
        "total_estimated_cost_usd": round(
            sum(float(result["estimated_cost_usd"]) for result in results), 4
        ),
        "anthropic_judge_cost_usd": round(
            sum(float(result.get("anthropic_judge_cost_usd") or 0.0) for result in results), 6
        ),
        "minimum_aggregate_pass_rate": 0.8,
        "minimum_per_format_pass_rate": 0.8,
        "per_task_pass_threshold": 0.70,
    }


def anthropic_spend_report(results: list[dict[str, Any]]) -> dict[str, Any]:
    calls = []
    for result in results:
        anthropic = result.get("quality", {}).get("anthropic", {})
        calls.append(
            {
                "task_id": result["id"],
                "format": result["format"],
                "model": anthropic.get("model"),
                "input_tokens": anthropic.get("input_tokens", 0),
                "output_tokens": anthropic.get("output_tokens", 0),
                "cost_usd": anthropic.get("cost_usd", 0.0),
                "latency_ms": anthropic.get("latency_ms", 0),
            }
        )
    return {
        "suite": "productivity-bench",
        "model": "claude-sonnet-4-6",
        "call_count": len(calls),
        "total_cost_usd": round(sum(float(call["cost_usd"] or 0.0) for call in calls), 6),
        "total_input_tokens": sum(int(call["input_tokens"] or 0) for call in calls),
        "total_output_tokens": sum(int(call["output_tokens"] or 0) for call in calls),
        "calls": calls,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Run the W7 productivity benchmark once per task.")
    parser.add_argument("--tasks", default=str(ROOT / "eval" / "productivity" / "tasks.yaml"))
    parser.add_argument("--output", default=str(ROOT / "outputs" / "v1_3" / "productivity-bench.json"))
    args = parser.parse_args()
    tasks = load_tasks(Path(args.tasks))
    results = [run_task(task, ROOT) for task in tasks]
    output = {
        "suite": "productivity-bench",
        "task_count": len(tasks),
        "summary": summarize(results),
        "results": results,
    }
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, indent=2), encoding="utf-8")
    (output_path.parent / "anthropic-spend.json").write_text(
        json.dumps(anthropic_spend_report(results), indent=2), encoding="utf-8"
    )
    print(json.dumps(output["summary"], indent=2))


if __name__ == "__main__":
    main()
