from __future__ import annotations

# ruff: noqa: E501,I001

import json
import sys
import tempfile
import urllib.request
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SKILLS_ROOT = Path(__file__).resolve().parents[2]
if str(SKILLS_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILLS_ROOT))

from shared.self_review import self_review  # noqa: E402
from shared.llm_content import generate_content  # noqa: E402

STYLE_FONTS = {
    "minimal": {"font": "Arial", "accent": colors.HexColor("#2E3440")},
    "academic": {"font": "Times New Roman", "accent": colors.HexColor("#4B5563")},
    "business": {"font": "Aptos", "accent": colors.HexColor("#1F4E79")},
}
TEMPLATE_DIR = Path(__file__).resolve().parents[1] / "templates"

DOCUMENT_PLANS = [
    (
        ("business memo", "phased rollout"),
        "business",
        ["docx", "pdf"],
        "Business Memo: Phased Automation Rollout",
        [
            ("Background", ["Leadership is considering a staged deployment of internal automation workflows."]),
            ("Recommendation", ["Approve a phased rollout with explicit success gates and review checkpoints."]),
            ("Rationale", ["A staged approach captures productivity gains while limiting operational risk."]),
            ("Risks and Mitigations", ["Key risks include quality drift, unclear ownership, and adoption friction."]),
            ("Decision Request", ["Approve the pilot scope, owners, and criteria for expansion."]),
        ],
    ),
    (
        ("academic", "mini report"),
        "academic",
        ["docx"],
        "Mini Report: AI Assistants and Knowledge Worker Productivity",
        [
            ("Abstract", ["This report summarizes how AI assistants can influence productivity outcomes."]),
            ("Method Summary", ["The analysis synthesizes workflow observations, adoption patterns, and quality concerns."]),
            ("Findings", ["Assistants help most when tasks have clear goals, reviewable outputs, and trusted controls."]),
            ("Limitations", ["Results vary by role maturity, data sensitivity, and measurement quality."]),
            ("Conclusion", ["Productivity impact depends on coupling speed with accountability and verification."]),
        ],
    ),
    (
        ("proposal", "pilot"),
        "business",
        ["pdf"],
        "Client Proposal: Secure Agent Workflow Pilot",
        [
            ("Scope", ["Deliver a four-week pilot focused on high-value, reviewable workflows."]),
            ("Deliverables", ["Provide configured workflows, operating guidance, and pilot reporting artifacts."]),
            ("Timeline", ["Use weekly milestones for discovery, build, validation, and handoff."]),
            ("Responsibilities", ["The client supplies owners and sample work; the delivery team manages implementation."]),
            ("Acceptance Criteria", ["The pilot succeeds when outputs meet quality, security, and cycle-time targets."]),
        ],
    ),
    (
        ("operating guide", "table of contents"),
        "minimal",
        ["docx"],
        "Operating Guide for Productivity Evaluations",
        [
            ("Setup", ["Prepare tools, credentials, task files, and output locations before execution."]),
            ("Task Selection", ["Choose prompts that require reasoning without pre-supplying the artifact outline."]),
            ("Execution", ["Run each benchmark task in a clean workspace and capture all generated artifacts."]),
            ("Review", ["Check structure, quality, and prompt alignment before recording pass rates."]),
            ("Troubleshooting", ["Diagnose missing files, malformed documents, and weak content separately."]),
            ("Governance", ["Preserve auditability with backups, costs, and reviewer notes."]),
            ("Reporting", ["Publish summary metrics, known limitations, and follow-up actions."]),
        ],
    ),
    (
        ("launch checklist",),
        "minimal",
        ["docx", "pdf"],
        "Benchmark Launch Checklist",
        [
            ("Preparation", ["Confirm tasks, rubrics, runner behavior, and artifact naming conventions."]),
            ("Quality Review", ["Verify generated files open cleanly and meet structural gates."]),
            ("Communications", ["Prepare release notes, stakeholder updates, and support guidance."]),
            ("Launch Day", ["Monitor runs, triage regressions, and keep rollback criteria visible."]),
            ("Post-Launch", ["Collect feedback, document issues, and schedule the next benchmark update."]),
        ],
    ),
    (
        ("meeting recap",),
        "business",
        ["docx"],
        "Meeting Recap: Productivity Benchmark Update",
        [
            ("Decisions", ["The team agreed to remove pre-supplied outlines and require prompt-derived artifacts."]),
            ("Action Items", ["Runner, skill behavior, QA, and release communication owners were assigned."]),
            ("Risks", ["The main risk is weakening quality if prompt interpretation is too shallow."]),
            ("Next Review", ["The team will review structural validation and sample outputs before release."]),
        ],
    ),
    (
        ("incident report",),
        "business",
        ["pdf"],
        "Incident Report: Over-Specified Productivity Benchmark Tasks",
        [
            ("Impact", ["Benchmark results overstated reasoning ability because outlines were supplied in task input."]),
            ("Root Cause", ["Task definitions embedded sections, bullets, and chart values that scripts could copy directly."]),
            ("Detection", ["Review identified that agents were formatting provided plans rather than creating them."]),
            ("Corrective Actions", ["Remove input outlines, preserve only structural gates, and back up the old suite."]),
            ("Prevention", ["Audit future tasks for hidden answer keys and require open-ended prompts."]),
        ],
    ),
    (
        ("research brief",),
        "academic",
        ["docx", "pdf"],
        "Research Brief: Useful Productivity Agents",
        [
            ("Question", ["What makes agents useful in real business workflows rather than impressive demos?"]),
            ("Evidence Themes", ["Teams value reliable completion, reviewable artifacts, and safe escalation paths."]),
            ("Interpretation", ["Usefulness depends on fitting into decision workflows and accountability structures."]),
            ("Caveats", ["Measurements can be distorted by toy prompts, hidden scaffolding, or narrow success gates."]),
            ("Evaluation Criteria", ["Assess autonomy, artifact quality, structural validity, and operational trust."]),
        ],
    ),
    (
        ("standard operating procedure", "sop"),
        "minimal",
        ["docx"],
        "Standard Operating Procedure: Artifact Review",
        [
            ("Purpose", ["Define a repeatable process for reviewing generated productivity artifacts."]),
            ("Scope", ["Apply this SOP to benchmark outputs before publication or release reporting."]),
            ("Roles", ["Assign an executor, structural reviewer, quality reviewer, and escalation owner."]),
            ("Procedure", ["Inspect file presence, openability, prompt alignment, and evidence of independent reasoning."]),
            ("Records", ["Store outputs, scores, reviewer notes, and any corrective actions."]),
            ("Escalation", ["Escalate security, data, or systemic quality failures immediately."]),
        ],
    ),
    (
        ("executive summary",),
        "business",
        ["pdf"],
        "Executive Summary: Productivity Benchmark Update",
        [
            ("Problem", ["The previous suite supplied outlines that reduced the benchmark's reasoning demands."]),
            ("Change", ["Tasks now rely on open-ended prompts and structural gates only."]),
            ("Benefits", ["The updated suite better measures planning, synthesis, and artifact construction."]),
            ("Risks", ["Prompt-derived outputs may vary, so quality judging and diagnostics remain important."]),
            ("Next Steps", ["Run the revised benchmark, inspect failures, and document the scoring impact."]),
        ],
    ),
]


def workspace_dir(payload: dict[str, Any]) -> Path:
    explicit = payload.get("output_dir")
    if explicit:
        out_dir = Path(str(explicit))
    else:
        task = str(payload.get("task") or payload.get("session_id") or "default").strip("/") or "default"
        out_dir = Path("/workspace") / task
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir


def _fallback_outline_from_prompt(payload: dict[str, Any]) -> dict[str, Any]:
    prompt = str(payload.get("prompt") or "").casefold()
    for keywords, style, formats, title, sections in DOCUMENT_PLANS:
        if any(keyword in prompt for keyword in keywords):
            payload.setdefault("style", style)
            payload.setdefault("formats", formats)
            return {
                "title": title,
                "sections": [
                    {"heading": heading, "paragraphs": paragraphs}
                    for heading, paragraphs in sections
                ],
            }
    payload.setdefault("style", "business")
    payload.setdefault("formats", ["docx", "pdf"])
    return {
        "title": "Productivity Artifact",
        "sections": [
            {"heading": "Context", "paragraphs": ["This document responds to the requested business context."]},
            {"heading": "Analysis", "paragraphs": ["The central considerations are quality, risk, ownership, and timing."]},
            {"heading": "Recommendation", "paragraphs": ["Proceed with clear success criteria and review checkpoints."]},
            {"heading": "Next Steps", "paragraphs": ["Assign owners, confirm milestones, and track outcomes."]},
        ],
    }


def _outline_from_prompt(payload: dict[str, Any]) -> dict[str, Any]:
    generated = generate_content(
        str(payload.get("prompt") or ""),
        "document",
        dict(payload.get("expected_schema") or payload.get("schema") or {}),
    )
    outline = _normalize_generated_outline(generated, payload)
    return outline or _fallback_outline_from_prompt(payload)


def _normalize_generated_outline(generated: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any] | None:
    raw_sections = generated.get("sections") or []
    if not isinstance(raw_sections, list) or not raw_sections:
        return None
    style = str(generated.get("style") or payload.get("style") or "business").lower()
    if style in STYLE_FONTS:
        payload.setdefault("style", style)
    formats = [str(item).lower() for item in generated.get("formats") or []]
    valid_formats = [item for item in formats if item in {"docx", "pdf"}]
    if valid_formats:
        payload.setdefault("formats", valid_formats)
    sections: list[dict[str, Any]] = []
    for index, section in enumerate(raw_sections, start=1):
        if not isinstance(section, dict):
            continue
        heading = str(section.get("heading") or section.get("title") or section.get("name") or f"Section {index}")
        paragraphs = section.get("paragraphs") or section.get("body") or section.get("content") or []
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        normalized: dict[str, Any] = {
            "heading": heading,
            "paragraphs": [str(item) for item in list(paragraphs) if str(item).strip()],
        }
        tables = section.get("tables")
        if isinstance(tables, list):
            normalized["tables"] = [table for table in tables if isinstance(table, dict)]
        images = section.get("images")
        if isinstance(images, list):
            normalized["images"] = images
        sections.append(normalized)
    if not sections:
        return None
    return {
        "title": str(generated.get("title") or payload.get("title") or "Document"),
        "sections": sections,
    }


def _image_path(ref: str, tmp_dir: Path) -> Path:
    if ref.startswith(("http://", "https://")):
        target = tmp_dir / (Path(ref.split("?", 1)[0]).name or "image")
        with urllib.request.urlopen(ref, timeout=15) as response:  # noqa: S310 - user-provided skill input
            target.write_bytes(response.read())
        return target
    return Path(ref)


def create_documents(payload: dict[str, Any]) -> list[Path]:
    formats = [str(item).lower() for item in payload.get("formats") or ["docx", "pdf"]]
    outline = dict(payload.get("outline") or payload)
    if not outline.get("sections"):
        outline = _outline_from_prompt(payload)
        formats = [str(item).lower() for item in payload.get("formats") or formats]
    out_dir = workspace_dir(payload)
    outputs: list[Path] = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        if "docx" in formats:
            outputs.append(
                _create_docx(out_dir / "output.docx", outline, str(payload.get("style") or "minimal"), tmp_dir)
            )
        if "pdf" in formats:
            outputs.append(
                _create_pdf(out_dir / "output.pdf", outline, str(payload.get("style") or "minimal"), tmp_dir)
            )
    return outputs


def _review_plan(payload: dict[str, Any]) -> dict[str, Any]:
    outline = payload.get("outline") or payload
    if isinstance(outline, dict) and outline.get("sections"):
        return dict(outline)
    return _outline_from_prompt(dict(payload))


def _create_docx(path: Path, outline: dict[str, Any], style_name: str, tmp_dir: Path) -> Path:
    template_path = TEMPLATE_DIR / f"{style_name}.docx"
    doc = Document(str(template_path)) if template_path.exists() else Document()
    _apply_docx_style(doc, style_name)
    if template_path.exists():
        doc.add_page_break()
    doc.add_heading(str(outline.get("title") or "Document"), 0)
    sections = list(outline.get("sections") or [])
    if len(sections) > 5:
        doc.add_heading("Table of Contents", 1)
        for idx, section in enumerate(sections, start=1):
            doc.add_paragraph(f"{idx}. {section.get('heading') or section.get('title')}")
        doc.add_page_break()
    for section in sections:
        doc.add_heading(str(section.get("heading") or section.get("title") or "Section"), 1)
        for paragraph in section.get("paragraphs") or section.get("body") or []:
            doc.add_paragraph(str(paragraph))
        for table in section.get("tables") or []:
            _add_docx_table(doc, table)
        for image in section.get("images") or []:
            image_path = _image_path(str(image), tmp_dir)
            if image_path.exists():
                doc.add_picture(str(image_path), width=Inches(5.8))
    doc.save(path)
    return path


def _apply_docx_style(doc: Document, style_name: str) -> None:
    font_name = STYLE_FONTS.get(style_name, STYLE_FONTS["minimal"])["font"]
    normal = doc.styles["Normal"]
    normal.font.name = str(font_name)
    normal.font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = str(font_name)
        style.font.bold = True
    if "Mantle Table" not in [style.name for style in doc.styles if style.type == WD_STYLE_TYPE.TABLE]:
        doc.styles.add_style("Mantle Table", WD_STYLE_TYPE.TABLE)


def _add_docx_table(doc: Document, spec: dict[str, Any]) -> None:
    headers = [str(item) for item in spec.get("headers") or []]
    rows = list(spec.get("rows") or [])
    table = doc.add_table(rows=1, cols=max(1, len(headers)))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        values = row if isinstance(row, list) else [row.get(header, "") for header in headers]
        for idx, value in enumerate(values[: len(cells)]):
            cells[idx].text = str(value)


def _create_pdf(path: Path, outline: dict[str, Any], style_name: str, tmp_dir: Path) -> Path:
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, title=str(outline.get("title") or "Document"))
    styles = getSampleStyleSheet()
    accent = STYLE_FONTS.get(style_name, STYLE_FONTS["minimal"])["accent"]
    story: list[Any] = [Paragraph(str(outline.get("title") or "Document"), styles["Title"]), Spacer(1, 18)]
    sections = list(outline.get("sections") or [])
    if len(sections) > 5:
        story.append(Paragraph("Table of Contents", styles["Heading1"]))
        for idx, section in enumerate(sections, start=1):
            heading = section.get("heading") or section.get("title")
            story.append(Paragraph(f"{idx}. {heading}", styles["BodyText"]))
        story.append(PageBreak())
    for section in sections:
        story.append(Paragraph(str(section.get("heading") or section.get("title") or "Section"), styles["Heading1"]))
        for paragraph in section.get("paragraphs") or section.get("body") or []:
            story.append(Paragraph(str(paragraph), styles["BodyText"]))
            story.append(Spacer(1, 8))
        for table in section.get("tables") or []:
            story.append(_pdf_table(table, accent))
            story.append(Spacer(1, 10))
        for image in section.get("images") or []:
            image_path = _image_path(str(image), tmp_dir)
            if image_path.exists():
                story.append(Image(str(image_path), width=360, height=220, kind="proportional"))
    doc.build(story)
    return path


def _pdf_table(spec: dict[str, Any], accent: Any) -> Table:
    headers = [str(item) for item in spec.get("headers") or []]
    rows = list(spec.get("rows") or [])
    data = [headers] if headers else []
    data.extend(row if isinstance(row, list) else [row.get(header, "") for header in headers] for row in rows)
    table = Table(data or [[""]])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), accent),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ]
        )
    )
    return table


def main() -> None:
    payload = json.loads(sys.stdin.read() or "{}")
    outputs = create_documents(payload)
    review_plan = _review_plan(payload)
    reviews = [
        {"artifact_id": path.name, **self_review(path, review_plan, "document", max_iterations=2)}
        for path in outputs
        if path.suffix == ".docx"
    ]
    print(
        json.dumps(
            {
                "paths": [str(path) for path in outputs],
                "formats": [path.suffix.lstrip(".") for path in outputs],
                "self_review": reviews[0] if len(reviews) == 1 else reviews,
            }
        )
    )


if __name__ == "__main__":
    main()
