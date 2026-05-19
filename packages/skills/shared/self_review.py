from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt

Issue = dict[str, Any]


def self_review(
    artifact_path: str | Path,
    plan: dict[str, Any],
    format_type: str,
    max_iterations: int = 2,
) -> dict[str, int]:
    path = Path(artifact_path)
    review_type = format_type.lower().lstrip(".")
    validators = {
        "slides": validate_slides,
        "pptx": validate_slides,
        "spreadsheet": validate_spreadsheet,
        "xlsx": validate_spreadsheet,
        "document": validate_document,
        "docx": validate_document,
    }
    fixers = {
        "slides": fix_slides,
        "pptx": fix_slides,
        "spreadsheet": fix_spreadsheet,
        "xlsx": fix_spreadsheet,
        "document": fix_document,
        "docx": fix_document,
    }
    if review_type not in validators:
        raise ValueError(f"Unsupported self-review format: {format_type}")

    validate = validators[review_type]
    fix = fixers[review_type]
    issues = validate(path, plan)
    initial_issue_count = len(issues)
    iterations_used = 0

    while issues and iterations_used < max_iterations:
        fix(path, issues, plan)
        iterations_used += 1
        issues = validate(path, plan)

    return {
        "iterations_used": iterations_used,
        "issues_found": initial_issue_count,
        "issues_resolved": max(0, initial_issue_count - len(issues)),
    }


def validate_slides(path: str | Path, plan: dict[str, Any]) -> list[Issue]:
    presentation = Presentation(path)
    slide_texts = [_slide_text(slide).casefold() for slide in presentation.slides]
    issues: list[Issue] = []
    for section in _planned_sections(plan):
        title = section["title"]
        if not any(title.casefold() in text for text in slide_texts):
            issues.append({"type": "missing_section", "section": title})
    return issues


def validate_spreadsheet(path: str | Path, plan: dict[str, Any]) -> list[Issue]:
    workbook = load_workbook(path, data_only=False)
    issues: list[Issue] = []
    for sheet_plan in _planned_sheets(plan):
        name = sheet_plan["name"]
        expected_headers = [str(header) for header in sheet_plan.get("headers") or []]
        if name not in workbook.sheetnames:
            issues.append({"type": "missing_sheet", "sheet": name, "headers": expected_headers})
            continue
        worksheet = workbook[name]
        actual_headers = [str(cell.value) for cell in worksheet[1] if cell.value not in (None, "")]
        if not actual_headers:
            issues.append({"type": "empty_sheet", "sheet": name, "headers": expected_headers})
        missing_headers = [header for header in expected_headers if header not in actual_headers]
        if missing_headers:
            issues.append({"type": "missing_headers", "sheet": name, "headers": missing_headers})
    return issues


def validate_document(path: str | Path, plan: dict[str, Any]) -> list[Issue]:
    document = Document(path)
    paragraphs = document.paragraphs
    issues: list[Issue] = []
    for section in _planned_sections(plan):
        title = section["title"]
        heading_index = next(
            (idx for idx, paragraph in enumerate(paragraphs) if paragraph.text.strip().casefold() == title.casefold()),
            None,
        )
        if heading_index is None:
            issues.append({"type": "missing_section", "section": title})
            continue
        body: list[str] = []
        for paragraph in paragraphs[heading_index + 1 :]:
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                break
            if paragraph.text.strip():
                body.append(paragraph.text.strip())
        if not body:
            issues.append({"type": "empty_section", "section": title})
    return issues


def fix_slides(path: str | Path, issues: list[Issue], plan: dict[str, Any]) -> None:
    presentation = Presentation(path)
    for issue in issues:
        if issue.get("type") != "missing_section":
            continue
        section = _find_section(plan, str(issue.get("section") or ""))
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.55), Inches(11.8), Inches(0.75))
        title_box.text_frame.text = str(issue["section"])
        title_box.text_frame.paragraphs[0].font.size = Pt(30)
        title_box.text_frame.paragraphs[0].font.bold = True
        bullets = _section_body_items(section) or ["Review and complete this planned section."]
        body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(10.8), Inches(4.8))
        text_frame = body_box.text_frame
        text_frame.clear()
        for idx, bullet in enumerate(bullets):
            paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
            paragraph.text = str(bullet)
            paragraph.font.size = Pt(20)
    presentation.save(path)


def fix_spreadsheet(path: str | Path, issues: list[Issue], plan: dict[str, Any]) -> None:
    workbook = load_workbook(path, data_only=False)
    planned = {sheet["name"]: sheet for sheet in _planned_sheets(plan)}
    for issue in issues:
        sheet_name = str(issue.get("sheet") or "Sheet")
        headers = list(issue.get("headers") or planned.get(sheet_name, {}).get("headers") or ["Item", "Value"])
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.append(headers)
            worksheet.append(["Self-review placeholder", "Added missing planned sheet"][: max(1, len(headers))])
            continue
        worksheet = workbook[sheet_name]
        actual_headers = [cell.value for cell in worksheet[1]] if worksheet.max_row else []
        if issue.get("type") == "empty_sheet" and headers:
            for index, header in enumerate(headers, start=1):
                worksheet.cell(row=1, column=index, value=header)
        else:
            for header in headers:
                if header not in actual_headers:
                    worksheet.cell(row=1, column=worksheet.max_column + 1, value=header)
    workbook.save(path)


def fix_document(path: str | Path, issues: list[Issue], plan: dict[str, Any]) -> None:
    document = Document(path)
    for issue in issues:
        if issue.get("type") not in {"missing_section", "empty_section"}:
            continue
        section = _find_section(plan, str(issue.get("section") or ""))
        if issue.get("type") == "missing_section":
            document.add_heading(str(issue["section"]), 1)
        for paragraph in _section_body_items(section) or ["Review and complete this planned section."]:
            document.add_paragraph(str(paragraph))
    document.save(path)


def _planned_sections(plan: dict[str, Any]) -> list[dict[str, Any]]:
    raw_sections = plan.get("sections") or plan.get("outline", {}).get("sections") or []
    sections: list[dict[str, Any]] = []
    for index, section in enumerate(raw_sections, start=1):
        if not isinstance(section, dict):
            continue
        title = str(section.get("title") or section.get("heading") or f"Section {index}").strip()
        if title:
            sections.append({**section, "title": title})
    return sections


def _planned_sheets(plan: dict[str, Any]) -> list[dict[str, Any]]:
    if plan.get("sheets"):
        return [_sheet_plan(sheet) for sheet in plan.get("sheets") or [] if isinstance(sheet, dict)]

    mode = str(plan.get("mode") or "table")
    sheets: list[dict[str, Any]]
    if mode == "financial_model":
        sheets = [
            {"name": "Assumptions", "headers": ["Assumption", "Value"]},
            {"name": "P&L", "headers": ["Line Item", "Year 1", "Year 2", "Year 3"]},
            {"name": "Balance Sheet", "headers": ["Line Item", "Year 1", "Year 2", "Year 3"]},
            {"name": "Cash Flow", "headers": ["Line Item", "Year 1", "Year 2", "Year 3"]},
        ]
    elif mode == "budget":
        sheets = [
            {"name": "Budget", "headers": ["Month", "Planned", "Actual", "Variance"]},
            {"name": "Summary", "headers": ["Metric", "Value"]},
        ]
    elif mode == "comparison_matrix":
        criteria = [str(criterion) for criterion in plan.get("criteria") or []]
        sheets = [{"name": "Comparison Matrix", "headers": ["Option", *criteria, "Total Score"]}]
    else:
        sheets = [{"name": "Data", "headers": _table_headers(plan)}]
    if plan.get("summary") and not any(sheet["name"] == "Summary" for sheet in sheets):
        sheets.append({"name": "Summary", "headers": ["Question", "Answer"]})
    if plan.get("charts"):
        sheets.append({"name": "Charts", "headers": []})
    return sheets


def _sheet_plan(sheet: dict[str, Any]) -> dict[str, Any]:
    name = str(sheet.get("name") or sheet.get("title") or "Sheet")
    headers = sheet.get("headers") or _schema_headers(sheet.get("schema") or [])
    return {"name": name, "headers": [str(header) for header in headers]}


def _table_headers(plan: dict[str, Any]) -> list[str]:
    headers = _schema_headers(plan.get("schema") or [])
    if str(plan.get("mode") or "") == "data_cleaning":
        headers = [_normalize_header(header) for header in headers]
    if headers:
        return headers
    keys = sorted({str(key) for row in plan.get("rows") or [] if isinstance(row, dict) for key in row})
    return keys


def _schema_headers(schema: list[Any]) -> list[str]:
    return [str(column.get("name")) for column in schema if isinstance(column, dict) and column.get("name")]


def _normalize_header(header: str) -> str:
    return header.strip().lower().replace(" ", "_").replace("-", "_")


def _find_section(plan: dict[str, Any], title: str) -> dict[str, Any]:
    for section in _planned_sections(plan):
        if section["title"].casefold() == title.casefold():
            return section
    return {"title": title}


def _section_body_items(section: dict[str, Any]) -> list[Any]:
    return list(
        section.get("bullets") or section.get("items") or section.get("paragraphs") or section.get("body") or []
    )


def _slide_text(slide: Any) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            parts.append(shape.text)
    return "\n".join(parts)
