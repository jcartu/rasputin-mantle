from __future__ import annotations

import json
import sys
import tempfile
import urllib.request
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TEMPLATES = {
    "minimal": {"accent": RGBColor(46, 52, 64), "bg": RGBColor(255, 255, 255), "fg": RGBColor(35, 35, 35)},
    "corporate": {"accent": RGBColor(31, 78, 121), "bg": RGBColor(245, 248, 252), "fg": RGBColor(20, 32, 48)},
    "pitch": {"accent": RGBColor(116, 60, 255), "bg": RGBColor(16, 18, 27), "fg": RGBColor(245, 245, 250)},
    "research": {"accent": RGBColor(17, 111, 93), "bg": RGBColor(250, 250, 245), "fg": RGBColor(28, 46, 43)},
}


def workspace_path(payload: dict[str, Any], filename: str) -> Path:
    explicit = payload.get("output_dir")
    if explicit:
        out_dir = Path(str(explicit))
    else:
        task = str(payload.get("task") or payload.get("session_id") or "default").strip("/") or "default"
        out_dir = Path("/workspace") / task
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / filename


def _set_background(slide: Any, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_footer(slide: Any, template: str, slide_number: int, style: dict[str, RGBColor]) -> None:
    box = slide.shapes.add_textbox(Inches(0.35), Inches(7.05), Inches(12.6), Inches(0.25))
    frame = box.text_frame
    frame.text = f"{template.title()} • {slide_number}"
    para = frame.paragraphs[0]
    para.font.size = Pt(8)
    para.font.color.rgb = style["accent"]


def _write_title(shape: Any, text: str, style: dict[str, RGBColor], size: int = 34) -> None:
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = style["fg"]


def _write_bullets(shape: Any, bullets: list[Any], style: dict[str, RGBColor]) -> None:
    tf = shape.text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        paragraph.text = str(bullet)
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = style["fg"]


def _image_path(ref: str, tmp_dir: Path) -> Path:
    if ref.startswith(("http://", "https://")):
        target = tmp_dir / (Path(ref.split("?", 1)[0]).name or "image")
        with urllib.request.urlopen(ref, timeout=15) as response:  # noqa: S310 - user-provided skill input
            target.write_bytes(response.read())
        return target
    return Path(ref)


def _add_images(slide: Any, refs: list[Any], tmp_dir: Path) -> None:
    x = Inches(8.2)
    y = Inches(1.55)
    for index, ref in enumerate(refs[:2]):
        path = _image_path(str(ref), tmp_dir)
        if path.exists():
            slide.shapes.add_picture(str(path), x, y + Inches(index * 2.4), width=Inches(4.4))


def _chart_image(chart: dict[str, Any], tmp_dir: Path) -> Path:
    labels = [str(item) for item in chart.get("labels", [])]
    values = [float(item) for item in chart.get("values", [])]
    kind = str(chart.get("type") or "bar")
    title = str(chart.get("title") or "")
    fig, ax = plt.subplots(figsize=(5, 3), dpi=160)
    if kind == "line":
        ax.plot(labels, values, marker="o")
    elif kind == "pie":
        ax.pie(values, labels=labels, autopct="%1.0f%%")
    else:
        ax.bar(labels, values)
    ax.set_title(title)
    if kind != "pie":
        ax.tick_params(axis="x", rotation=25)
        ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    target = tmp_dir / "chart.png"
    fig.savefig(target, transparent=False)
    plt.close(fig)
    return target


def _add_notes(slide: Any, notes: str) -> None:
    if not notes:
        return
    notes_slide = getattr(slide, "notes_slide", None)
    if notes_slide is not None:
        text_frame = notes_slide.notes_text_frame
        text_frame.text = notes
        return
    box = slide.shapes.add_textbox(Inches(0.35), Inches(6.72), Inches(12.4), Inches(0.25))
    box.name = "Speaker Notes"
    box.text_frame.text = f"Notes: {notes}"
    box.text_frame.paragraphs[0].font.size = Pt(7)


def create_presentation(payload: dict[str, Any]) -> Path:
    template = str(payload.get("template") or "minimal").lower()
    style = TEMPLATES.get(template, TEMPLATES["minimal"])
    outline = payload.get("outline") or payload
    sections = list(outline.get("sections") or [])
    title = str(outline.get("title") or payload.get("title") or "Presentation")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        title_slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(title_slide, style["bg"])
        title_box = title_slide.shapes.add_textbox(Inches(0.8), Inches(2.25), Inches(11.8), Inches(1.1))
        _write_title(title_box, title, style, 44)
        subtitle = str(outline.get("subtitle") or "Generated by Mantle Slides")
        box = title_slide.shapes.add_textbox(Inches(0.85), Inches(3.45), Inches(10.8), Inches(0.6))
        box.text_frame.text = subtitle
        box.text_frame.paragraphs[0].font.size = Pt(20)
        box.text_frame.paragraphs[0].font.color.rgb = style["accent"]
        _add_footer(title_slide, template, 1, style)
        _add_notes(title_slide, str(outline.get("notes") or "Introduce the deck purpose and agenda."))

        for idx, section in enumerate(sections, start=2):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            _set_background(slide, style["bg"])
            section_title = str(section.get("title") or f"Section {idx - 1}")
            section_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(12.1), Inches(0.7))
            _write_title(section_box, section_title, style, 30)
            bullets = list(section.get("bullets") or section.get("items") or [])
            bullets_box = slide.shapes.add_textbox(Inches(0.85), Inches(1.45), Inches(7.0), Inches(4.9))
            _write_bullets(bullets_box, bullets, style)
            if section.get("images"):
                _add_images(slide, list(section.get("images") or []), tmp_dir)
            if section.get("chart"):
                chart_path = _chart_image(dict(section.get("chart") or {}), tmp_dir)
                slide.shapes.add_picture(str(chart_path), Inches(8.1), Inches(1.5), width=Inches(4.5))
            _add_footer(slide, template, idx, style)
            _add_notes(slide, str(section.get("notes") or ""))

        output = workspace_path(payload, "slides.pptx")
        prs.save(output)
        return output


def main() -> None:
    payload = json.loads(sys.stdin.read() or "{}")
    output = create_presentation(payload)
    print(
        json.dumps(
            {
                "path": str(output),
                "template": payload.get("template", "minimal"),
                "slide_count": len(Presentation(output).slides),
            }
        )
    )


if __name__ == "__main__":
    main()
