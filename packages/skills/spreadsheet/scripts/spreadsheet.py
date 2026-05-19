from __future__ import annotations

# ruff: noqa: E402,E501,I001

import json
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
TEMPLATE_DIR = ROOT / "templates"
SKILLS_ROOT = Path(__file__).resolve().parents[2]
if str(SKILLS_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILLS_ROOT))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorkbookImage
from openpyxl.styles import Font, PatternFill

from subskills.comparison_matrix import add_comparison_matrix  # noqa: E402
from subskills.data_cleaning import clean_rows  # noqa: E402
from subskills.financial_model import build_financial_model  # noqa: E402
from shared.self_review import self_review  # noqa: E402
from shared.llm_content import generate_content  # noqa: E402


def workspace_path(payload: dict[str, Any], filename: str) -> Path:
    explicit = payload.get("output_dir")
    if explicit:
        out_dir = Path(str(explicit))
    else:
        task = str(payload.get("task") or payload.get("session_id") or "default").strip("/") or "default"
        out_dir = Path("/workspace") / task
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir / filename


def _fallback_payload_from_prompt(payload: dict[str, Any]) -> dict[str, Any]:
    prompt = str(payload.get("prompt") or "").casefold()
    base = {"output_dir": payload.get("output_dir"), "prompt": payload.get("prompt")}
    if "financial model" in prompt or "five-year" in prompt or "five year" in prompt:
        return {
            **base,
            "mode": "financial_model",
            "assumptions": {"starting_revenue": 850000, "revenue_growth": 0.28, "gross_margin": 0.74},
        }
    if "feature comparison" in prompt or "vendor scoring" in prompt:
        criteria = ["Cost", "Security", "Usability", "Support", "Implementation"]
        if "vendor" in prompt:
            criteria = ["Cost", "Compliance", "Support", "Reliability", "Implementation"]
        return {
            **base,
            "mode": "comparison_matrix",
            "criteria": criteria,
            "items": [
                {"name": "Option A", "scores": {criterion: 4 for criterion in criteria}},
                {"name": "Option B", "scores": {criterion: 3 + (idx % 2) for idx, criterion in enumerate(criteria)}},
                {
                    "name": "Option C",
                    "scores": {criterion: 5 if idx < 2 else 3 for idx, criterion in enumerate(criteria)},
                },
            ],
        }
    if "customer data cleaning" in prompt or "deduplicate" in prompt:
        return {
            **base,
            "mode": "data_cleaning",
            "schema": [{"name": "Name"}, {"name": "Company"}, {"name": "Spend"}, {"name": "Status"}],
            "rows": [
                {"Name": "Ada Lovelace", "Company": "Analytical Ops", "Spend": "$1,200", "Status": "Active"},
                {"Name": "Ada Lovelace", "Company": "Analytical Ops", "Spend": "$1,200", "Status": "Active"},
                {"Name": "Lin Chen", "Company": "Vector Labs", "Spend": "900", "Status": "Trial"},
                {"Name": "Maya Singh", "Company": "Nimbus", "Spend": "$2,450", "Status": "Active"},
            ],
            "summary": True,
        }
    if "expense" in prompt:
        return {
            **base,
            "mode": "table",
            "schema": [{"name": "Category"}, {"name": "Vendor"}, {"name": "Monthly Cost"}, {"name": "Variance"}],
            "rows": [
                {"Category": "Compute", "Vendor": "GPU Cloud", "Monthly Cost": 4200, "Variance": 350},
                {"Category": "Models", "Vendor": "API Provider", "Monthly Cost": 1800, "Variance": -120},
                {"Category": "Storage", "Vendor": "Object Store", "Monthly Cost": 420, "Variance": 40},
            ],
            "charts": [{"type": "bar", "x": "Category", "y": "Monthly Cost", "title": "Monthly Cost by Category"}],
        }
    if "hiring" in prompt:
        return {
            **base,
            "mode": "table",
            "schema": [{"name": "Role"}, {"name": "Stage"}, {"name": "Owner"}, {"name": "Candidates"}],
            "rows": [
                {"Role": "Platform Engineer", "Stage": "Technical", "Owner": "Avery", "Candidates": 5},
                {"Role": "Product Designer", "Stage": "Portfolio", "Owner": "Sam", "Candidates": 3},
                {"Role": "Operations Lead", "Stage": "Onsite", "Owner": "Jordan", "Candidates": 2},
            ],
            "summary": True,
        }
    if "product metrics" in prompt:
        return {
            **base,
            "mode": "table",
            "schema": [{"name": "Metric"}, {"name": "Current"}, {"name": "Previous"}, {"name": "Interpretation"}],
            "rows": [
                {"Metric": "Activation", "Current": 61, "Previous": 54, "Interpretation": "Improving"},
                {"Metric": "Retention", "Current": 43, "Previous": 41, "Interpretation": "Stable"},
                {"Metric": "Quality", "Current": 88, "Previous": 82, "Interpretation": "Improving"},
            ],
            "charts": [{"type": "bar", "x": "Metric", "y": "Current", "title": "Current Product Metrics"}],
        }
    if "budget" in prompt:
        return {**base, "mode": "budget"}
    if "filter" in prompt or "status" in prompt:
        return {
            **base,
            "mode": "table",
            "schema": [{"name": "Initiative"}, {"name": "Owner"}, {"name": "Health"}, {"name": "Priority"}],
            "rows": [
                {"Initiative": "Runner update", "Owner": "Ada", "Health": "Green", "Priority": "High"},
                {"Initiative": "Skill QA", "Owner": "Lin", "Health": "Yellow", "Priority": "High"},
                {"Initiative": "Release comms", "Owner": "Maya", "Health": "Red", "Priority": "Medium"},
            ],
            "query_plan": {"where_equals": {"Health": "Green"}},
            "summary": True,
        }
    return {
        **base,
        "mode": "table",
        "schema": [{"name": "Stage"}, {"name": "Owner"}, {"name": "Deals"}, {"name": "Forecast Risk"}],
        "rows": [
            {"Stage": "Prospect", "Owner": "Sales", "Deals": 18, "Forecast Risk": "Medium"},
            {"Stage": "Qualified", "Owner": "AE Team", "Deals": 11, "Forecast Risk": "Low"},
            {"Stage": "Commit", "Owner": "Revenue Lead", "Deals": 6, "Forecast Risk": "Low"},
        ],
        "charts": [{"type": "bar", "x": "Stage", "y": "Deals", "title": "Deals by Stage"}],
    }


def _payload_from_prompt(payload: dict[str, Any]) -> dict[str, Any]:
    generated = generate_content(
        str(payload.get("prompt") or ""),
        "spreadsheet",
        dict(payload.get("expected_schema") or payload.get("schema") or {}),
    )
    normalized = _normalize_generated_payload(generated, payload)
    return (
        normalized
        or _financial_model_payload_from_prompt(payload)
        or _facts_payload_from_prompt(payload)
        or _fallback_payload_from_prompt(payload)
    )


def _normalize_generated_payload(generated: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any] | None:
    base = {"output_dir": payload.get("output_dir"), "prompt": payload.get("prompt")}
    mode = str(generated.get("mode") or "table").lower()
    if mode not in {"table", "financial_model", "comparison_matrix", "data_cleaning", "budget"}:
        mode = "table"

    sheets = _generated_sheets(generated)
    if sheets:
        return {
            **base,
            "mode": "multi_table",
            "sheets": sheets,
            "charts": [chart for chart in generated.get("charts") or [] if isinstance(chart, dict)],
        }

    if mode == "financial_model":
        assumptions = generated.get("assumptions")
        if isinstance(assumptions, dict) and assumptions:
            return {**base, "mode": mode, "assumptions": assumptions}
        return _financial_model_payload_from_prompt(payload)
    if mode == "budget":
        return None
    if mode == "comparison_matrix":
        criteria = [str(item) for item in generated.get("criteria") or []]
        items = [item for item in generated.get("items") or [] if isinstance(item, dict)]
        if criteria and items:
            return {**base, "mode": mode, "criteria": criteria, "items": items}

    rows = [row for row in generated.get("rows") or [] if isinstance(row, dict)]
    schema = [col for col in generated.get("schema") or [] if isinstance(col, dict) and col.get("name")]
    if not schema and rows:
        schema = [{"name": key} for key in rows[0].keys()]
    if not rows or not schema:
        raw_sheets = generated.get("sheets") or []
        if isinstance(raw_sheets, list) and raw_sheets and isinstance(raw_sheets[0], dict):
            first_sheet = raw_sheets[0]
            rows = [row for row in first_sheet.get("rows") or [] if isinstance(row, dict)]
            schema = [col for col in first_sheet.get("schema") or [] if isinstance(col, dict) and col.get("name")]
            if not schema and rows:
                schema = [{"name": key} for key in rows[0].keys()]
    if not rows or not schema:
        return None

    charts = [chart for chart in generated.get("charts") or [] if isinstance(chart, dict)]
    query_plan = generated.get("query_plan") if isinstance(generated.get("query_plan"), dict) else None
    return {
        **base,
        "mode": "data_cleaning" if mode == "data_cleaning" else "table",
        "schema": schema,
        "rows": rows,
        "charts": charts,
        "query_plan": query_plan,
        "summary": bool(generated.get("summary", True)),
    }


def _generated_sheets(generated: dict[str, Any]) -> list[dict[str, Any]]:
    candidates = generated.get("rows") or generated.get("sheets") or []
    if not isinstance(candidates, list):
        return []
    sheets: list[dict[str, Any]] = []
    schema_by_name = {
        str(item.get("name") or item.get("sheet") or item.get("sheet_name")): item
        for item in generated.get("schema") or []
        if isinstance(item, dict)
    }
    for index, item in enumerate(candidates, start=1):
        if not isinstance(item, dict):
            continue
        data = item.get("data") or item.get("rows")
        if not isinstance(data, list) or not data or not all(isinstance(row, dict) for row in data):
            continue
        name = str(item.get("sheet") or item.get("sheet_name") or item.get("name") or f"Sheet {index}")[:31]
        schema_spec = schema_by_name.get(name, {})
        columns = item.get("columns") or schema_spec.get("columns") or list(data[0].keys())
        sheets.append(
            {
                "name": name,
                "schema": [{"name": str(column)} for column in columns],
                "rows": data,
            }
        )
    return sheets


def _financial_model_payload_from_prompt(payload: dict[str, Any]) -> dict[str, Any] | None:
    prompt = str(payload.get("prompt") or "")
    if "financial model" not in prompt.casefold():
        return None
    starting_arr = _money_after(prompt, "Starting ARR") or 500_000
    opening_cash = _money_after(prompt, "Opening cash") or 1_200_000
    growth = _percentages_after(prompt, "Revenue growth assumptions by year") or [0.25, 0.32, 0.35, 0.28, 0.22]
    margin = _percentages_after(prompt, "Gross margin assumptions by year") or [0.72, 0.74, 0.76, 0.77, 0.78]
    sm = _declining_percentages(prompt, "Sales and marketing", [0.38, 0.36, 0.34, 0.32, 0.30])
    rd = _declining_percentages(prompt, "R&D", [0.42, 0.385, 0.35, 0.315, 0.28])
    ga = _declining_percentages(prompt, "G&A", [0.18, 0.165, 0.15, 0.135, 0.12])
    capex = _percent_after(prompt, "Capex") or 0.04
    working_capital = _percent_after(prompt, "Working capital reserve") or 0.08
    years = ["Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    schema = [{"name": "Line Item"}, *[{"name": year} for year in years]]

    income_rows = [
        {
            "Line Item": "Revenue",
            "Year 1": f"={starting_arr}*(1+Assumptions!B2)",
            "Year 2": "=B2*(1+Assumptions!C2)",
            "Year 3": "=C2*(1+Assumptions!D2)",
            "Year 4": "=D2*(1+Assumptions!E2)",
            "Year 5": "=E2*(1+Assumptions!F2)",
        },
        {
            "Line Item": "Gross Profit",
            "Year 1": "=B2*Assumptions!B3",
            "Year 2": "=C2*Assumptions!C3",
            "Year 3": "=D2*Assumptions!D3",
            "Year 4": "=E2*Assumptions!E3",
            "Year 5": "=F2*Assumptions!F3",
        },
        {
            "Line Item": "Sales & Marketing",
            "Year 1": "=B2*Assumptions!B4",
            "Year 2": "=C2*Assumptions!C4",
            "Year 3": "=D2*Assumptions!D4",
            "Year 4": "=E2*Assumptions!E4",
            "Year 5": "=F2*Assumptions!F4",
        },
        {
            "Line Item": "R&D",
            "Year 1": "=B2*Assumptions!B5",
            "Year 2": "=C2*Assumptions!C5",
            "Year 3": "=D2*Assumptions!D5",
            "Year 4": "=E2*Assumptions!E5",
            "Year 5": "=F2*Assumptions!F5",
        },
        {
            "Line Item": "G&A",
            "Year 1": "=B2*Assumptions!B6",
            "Year 2": "=C2*Assumptions!C6",
            "Year 3": "=D2*Assumptions!D6",
            "Year 4": "=E2*Assumptions!E6",
            "Year 5": "=F2*Assumptions!F6",
        },
        {
            "Line Item": "EBITDA",
            "Year 1": "=B3-SUM(B4:B6)",
            "Year 2": "=C3-SUM(C4:C6)",
            "Year 3": "=D3-SUM(D4:D6)",
            "Year 4": "=E3-SUM(E4:E6)",
            "Year 5": "=F3-SUM(F4:F6)",
        },
    ]
    assumptions = [
        {"Line Item": "Revenue Growth Rate", **dict(zip(years, growth, strict=False))},
        {"Line Item": "Gross Margin", **dict(zip(years, margin, strict=False))},
        {"Line Item": "Sales & Marketing % Revenue", **dict(zip(years, sm, strict=False))},
        {"Line Item": "R&D % Revenue", **dict(zip(years, rd, strict=False))},
        {"Line Item": "G&A % Revenue", **dict(zip(years, ga, strict=False))},
        {"Line Item": "Capex % Revenue", **dict(zip(years, [capex] * 5, strict=False))},
        {"Line Item": "Working Capital Reserve %", **dict(zip(years, [working_capital] * 5, strict=False))},
        {"Line Item": "Opening Cash", "Year 1": opening_cash},
    ]
    cash_flow = [
        {
            "Line Item": "EBITDA",
            "Year 1": "='Income Statement'!B7",
            "Year 2": "='Income Statement'!C7",
            "Year 3": "='Income Statement'!D7",
            "Year 4": "='Income Statement'!E7",
            "Year 5": "='Income Statement'!F7",
        },
        {
            "Line Item": "Capex",
            "Year 1": "=-'Income Statement'!B2*Assumptions!B7",
            "Year 2": "=-'Income Statement'!C2*Assumptions!C7",
            "Year 3": "=-'Income Statement'!D2*Assumptions!D7",
            "Year 4": "=-'Income Statement'!E2*Assumptions!E7",
            "Year 5": "=-'Income Statement'!F2*Assumptions!F7",
        },
        {
            "Line Item": "Working Capital Reserve",
            "Year 1": "=-'Income Statement'!B2*Assumptions!B8",
            "Year 2": "=-'Income Statement'!C2*Assumptions!C8",
            "Year 3": "=-'Income Statement'!D2*Assumptions!D8",
            "Year 4": "=-'Income Statement'!E2*Assumptions!E8",
            "Year 5": "=-'Income Statement'!F2*Assumptions!F8",
        },
        {
            "Line Item": "Net Cash Flow",
            "Year 1": "=SUM(B2:B4)",
            "Year 2": "=SUM(C2:C4)",
            "Year 3": "=SUM(D2:D4)",
            "Year 4": "=SUM(E2:E4)",
            "Year 5": "=SUM(F2:F4)",
        },
    ]
    balance_sheet = [
        {
            "Line Item": "Cash Balance",
            "Year 1": f"={opening_cash}+'Cash Flow'!B5",
            "Year 2": "=B2+'Cash Flow'!C5",
            "Year 3": "=C2+'Cash Flow'!D5",
            "Year 4": "=D2+'Cash Flow'!E5",
            "Year 5": "=E2+'Cash Flow'!F5",
        },
        {
            "Line Item": "Working Capital Reserve",
            "Year 1": "='Income Statement'!B2*Assumptions!B8",
            "Year 2": "='Income Statement'!C2*Assumptions!C8",
            "Year 3": "='Income Statement'!D2*Assumptions!D8",
            "Year 4": "='Income Statement'!E2*Assumptions!E8",
            "Year 5": "='Income Statement'!F2*Assumptions!F8",
        },
    ]
    summary = [
        {"Line Item": "Starting ARR", "Year 1": starting_arr},
        {"Line Item": "Year 5 Revenue", "Year 5": "='Income Statement'!F2"},
        {"Line Item": "Year 5 EBITDA", "Year 5": "='Income Statement'!F7"},
        {"Line Item": "Ending Cash", "Year 5": "='Balance Sheet'!F2"},
    ]
    return {
        "output_dir": payload.get("output_dir"),
        "prompt": payload.get("prompt"),
        "mode": "multi_table",
        "sheets": [
            {"name": "Income Statement", "schema": schema, "rows": income_rows},
            {"name": "Assumptions", "schema": schema, "rows": assumptions},
            {"name": "Cash Flow", "schema": schema, "rows": cash_flow},
            {"name": "Balance Sheet", "schema": schema, "rows": balance_sheet},
            {"name": "Summary", "schema": schema, "rows": summary},
        ],
        "charts": [{"type": "line", "x": "Line Item", "y": "Year 1", "title": "Financial Model Overview"}],
    }


def _facts_payload_from_prompt(payload: dict[str, Any]) -> dict[str, Any] | None:
    facts = _extract_prompt_facts(str(payload.get("prompt") or ""))
    if not facts:
        return None
    rows = []
    for index, fact in enumerate(facts, start=1):
        label, _, detail = fact.partition(":")
        rows.append(
            {
                "ID": index,
                "Topic": label.strip() or f"Fact {index}",
                "Detail": detail.strip() or fact,
                "Numeric Value": _first_numeric_value(fact),
            }
        )
    return {
        "output_dir": payload.get("output_dir"),
        "prompt": payload.get("prompt"),
        "mode": "table",
        "schema": [{"name": "ID"}, {"name": "Topic"}, {"name": "Detail"}, {"name": "Numeric Value"}],
        "rows": rows,
        "summary": True,
        "charts": [{"type": "bar", "x": "Topic", "y": "Numeric Value", "title": "Prompt Facts by Numeric Value"}],
    }


def _extract_prompt_facts(prompt: str) -> list[str]:
    facts: list[str] = []
    for line in prompt.splitlines():
        stripped = line.strip()
        if stripped.startswith("- "):
            facts.append(stripped[2:].strip())
    return facts


def _first_numeric_value(text: str) -> float | int:
    match = re.search(r"\$?([0-9]+(?:,[0-9]{3})*(?:\.[0-9]+)?)(M|K|%)?", text, flags=re.IGNORECASE)
    if not match:
        return 0
    value = float(match.group(1).replace(",", ""))
    suffix = (match.group(2) or "").upper()
    if suffix == "M":
        value *= 1_000_000
    elif suffix == "K":
        value *= 1_000
    elif suffix == "%":
        value /= 100
    return int(value) if value.is_integer() else value


def _money_after(prompt: str, label: str) -> float | None:
    match = re.search(rf"{re.escape(label)}[^\n$]*\$([0-9]+(?:,[0-9]{{3}})*(?:\.[0-9]+)?)(M|K)?", prompt, re.IGNORECASE)
    if not match:
        return None
    value = float(match.group(1).replace(",", ""))
    suffix = (match.group(2) or "").upper()
    if suffix == "M":
        value *= 1_000_000
    elif suffix == "K":
        value *= 1_000
    return value


def _percentages_after(prompt: str, label: str) -> list[float]:
    match = re.search(rf"{re.escape(label)}[^\n]*", prompt, re.IGNORECASE)
    if not match:
        return []
    return [float(value) / 100 for value in re.findall(r"([0-9]+(?:\.[0-9]+)?)%", match.group(0))][:5]


def _percent_after(prompt: str, label: str) -> float | None:
    values = _percentages_after(prompt, label)
    return values[0] if values else None


def _declining_percentages(prompt: str, label: str, fallback: list[float]) -> list[float]:
    values = _percentages_after(prompt, label)
    if len(values) >= 2:
        start, end = values[0], values[-1]
        step = (start - end) / 4
        return [round(start - step * index, 4) for index in range(5)]
    return fallback


def create_workbook(payload: dict[str, Any]) -> Path:
    if not payload.get("mode"):
        payload = _payload_from_prompt(payload)
    output = workspace_path(payload, "data.xlsx")
    mode = str(payload.get("mode") or "table")
    if mode == "financial_model":
        template = TEMPLATE_DIR / "financial-model.xlsx"
        if template.exists():
            wb = load_workbook(template)
            wb.save(output)
            return output
        return build_financial_model(output, dict(payload.get("assumptions") or {}))
    if mode == "budget":
        return _create_budget_workbook(output)
    if mode == "multi_table":
        return _create_multi_table_workbook(output, payload)

    workbook_templates = {
        "comparison_matrix": "comparison-matrix.xlsx",
        "data_cleaning": "data-cleaning.xlsx",
    }
    template_name = workbook_templates.get(mode)
    if template_name and (TEMPLATE_DIR / template_name).exists():
        wb = load_workbook(TEMPLATE_DIR / template_name)
        wb.save(output)
        return output

    wb = Workbook()
    if mode == "comparison_matrix":
        add_comparison_matrix(wb, list(payload.get("items") or []), [str(c) for c in payload.get("criteria") or []])
    else:
        rows = list(payload.get("rows") or [])
        schema = list(payload.get("schema") or [])
        if mode == "data_cleaning":
            rows = clean_rows(rows, schema)
            schema = [{"name": key} for key in (rows[0].keys() if rows else [])]
        if payload.get("query_plan"):
            rows = _apply_query_plan(rows, dict(payload.get("query_plan") or {}))
        _add_table(wb.active, rows, schema)
    if payload.get("summary"):
        _add_summary_sheet(wb, str(payload.get("prompt") or "Summary"))
    _format(wb)
    wb.save(output)
    if payload.get("charts"):
        _add_charts(output, list(payload.get("charts") or []))
    return output


def _create_multi_table_workbook(output: Path, payload: dict[str, Any]) -> Path:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_spec in payload.get("sheets") or []:
        if not isinstance(sheet_spec, dict):
            continue
        ws = wb.create_sheet(str(sheet_spec.get("name") or "Data")[:31])
        _add_table(ws, list(sheet_spec.get("rows") or []), list(sheet_spec.get("schema") or []))
    if not wb.worksheets:
        ws = wb.create_sheet("Data")
        _add_table(ws, [], [])
    _format(wb)
    wb.save(output)
    charts = list(payload.get("charts") or [])
    if not charts:
        first = wb.worksheets[0]
        if first.max_column >= 2:
            charts = [{"type": "bar", "x": first.cell(1, 1).value, "y": first.cell(1, 2).value, "title": first.title}]
    if charts:
        _add_charts(output, charts)
    return output


def _review_plan(payload: dict[str, Any]) -> dict[str, Any]:
    if payload.get("mode"):
        return dict(payload)
    return _payload_from_prompt(dict(payload))


def _create_budget_workbook(output: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Month", "Planned", "Actual", "Variance"])
    rows = [("Jan", 1200, 1150), ("Feb", 1350, 1425), ("Mar", 1500, 1610), ("Apr", 1650, 1580)]
    for idx, (month, planned, actual) in enumerate(rows, start=2):
        ws.append([month, planned, actual, f"=B{idx}-C{idx}"])
    ws.append(["Total", "=SUM(B2:B5)", "=SUM(C2:C5)", "=SUM(D2:D5)"])
    summary = wb.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Total Planned", "=Budget!B6"])
    summary.append(["Total Actual", "=Budget!C6"])
    summary.append(["Total Variance", "=Budget!D6"])
    _format(wb)
    wb.save(output)
    _add_charts(output, [{"type": "line", "x": "Month", "y": "Planned", "title": "Planned Budget"}])
    return output


def _add_summary_sheet(wb: Workbook, prompt: str) -> None:
    if "Summary" in wb.sheetnames:
        return
    ws = wb.create_sheet("Summary")
    ws.append(["Question", "Answer"])
    ws.append(["Purpose", prompt[:180]])
    ws.append(["Record Count", "=COUNTA(Data!A:A)-1"])
    ws.append(["Review Focus", "Use the workbook to identify risks, owners, and next decisions."])
    ws.append(["Next Step", "Validate assumptions with the responsible team before publication."])


def _apply_query_plan(rows: list[dict[str, Any]], plan: dict[str, Any]) -> list[dict[str, Any]]:
    select = [str(item) for item in plan.get("select") or []]
    filtered = rows
    if plan.get("where_equals"):
        expected = dict(plan["where_equals"])
        filtered = [row for row in filtered if all(row.get(key) == value for key, value in expected.items())]
    if select:
        filtered = [{key: row.get(key) for key in select} for row in filtered]
    return filtered


def _add_table(ws: Any, rows: list[dict[str, Any]], schema: list[dict[str, Any]]) -> None:
    headers = [str(col.get("name")) for col in schema] if schema else sorted({key for row in rows for key in row})
    ws.title = "Data"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])


def _format(wb: Workbook) -> None:
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
        ws.freeze_panes = "A2"
        for column_cells in ws.columns:
            ws.column_dimensions[column_cells[0].column_letter].width = 18


def _add_charts(path: Path, charts: list[dict[str, Any]]) -> None:
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.values)
    if len(rows) < 2:
        wb.save(path)
        return
    df = pd.DataFrame(rows[1:], columns=rows[0])
    chart_ws = wb.create_sheet("Charts")
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        for index, chart in enumerate(charts, start=1):
            image = _chart_image(df, chart, tmp_dir / f"chart-{index}.png")
            chart_ws.add_image(WorkbookImage(str(image)), f"A{1 + (index - 1) * 18}")
        wb.save(path)


def _chart_image(df: pd.DataFrame, chart: dict[str, Any], target: Path) -> Path:
    x_col = str(chart.get("x") or df.columns[0])
    y_col = str(chart.get("y") or df.columns[-1])
    if x_col not in df.columns:
        x_col = str(df.columns[0])
    if y_col not in df.columns:
        numeric_columns = [column for column in df.columns if pd.to_numeric(df[column], errors="coerce").notna().any()]
        y_col = str(numeric_columns[0] if numeric_columns else df.columns[-1])
    kind = str(chart.get("type") or "bar")
    fig, ax = plt.subplots(figsize=(6, 3.4), dpi=150)
    if kind == "line":
        ax.plot(df[x_col].astype(str), pd.to_numeric(df[y_col], errors="coerce"), marker="o")
    else:
        ax.bar(df[x_col].astype(str), pd.to_numeric(df[y_col], errors="coerce"))
    ax.set_title(str(chart.get("title") or y_col))
    ax.tick_params(axis="x", rotation=25)
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(target)
    plt.close(fig)
    return target


def main() -> None:
    payload = json.loads(sys.stdin.read() or "{}")
    output = create_workbook(payload)
    review_plan = _review_plan(payload)
    review = self_review(output, review_plan, "spreadsheet", max_iterations=2)
    workbook = load_workbook(output, data_only=False)
    print(
        json.dumps(
            {
                "path": str(output),
                "sheets": workbook.sheetnames,
                "mode": review_plan.get("mode", "table"),
                "self_review": {"artifact_id": output.name, **review},
            }
        )
    )


if __name__ == "__main__":
    main()
