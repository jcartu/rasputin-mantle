from __future__ import annotations

from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def add_comparison_matrix(wb: Workbook, items: list[dict[str, Any]], criteria: list[str]) -> None:
    ws = wb.active
    ws.title = "Comparison Matrix"
    ws.append(["Option", *criteria, "Total Score"])
    for item in items:
        row = [item.get("name") or item.get("option") or "Option"]
        scores = item.get("scores") or item
        for criterion in criteria:
            row.append(scores.get(criterion, ""))
        row.append(f"=SUM(B{ws.max_row + 1}:{chr(65 + len(criteria))}{ws.max_row + 1})")
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="743CFF")
    ws.freeze_panes = "B2"
