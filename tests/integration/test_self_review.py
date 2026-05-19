from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from shared import self_review as sr

SLIDE_PLAN = {
    "sections": [
        {"title": "Context", "bullets": ["Why it matters"]},
        {"title": "Decision", "bullets": ["Approve"]},
    ]
}
SHEET_PLAN = {
    "sheets": [
        {"name": "Data", "headers": ["Metric", "Value"]},
        {"name": "Summary", "headers": ["Question", "Answer"]},
    ]
}
DOC_PLAN = {
    "sections": [
        {"heading": "Context", "paragraphs": ["Background"]},
        {"heading": "Decision", "paragraphs": ["Approve rollout."]},
    ]
}


def test_validators_accept_known_good_artifacts(tmp_path: Path) -> None:
    slides = _slides(tmp_path / "good.pptx", ["Context", "Decision"])
    spreadsheet = _spreadsheet(tmp_path / "good.xlsx", include_summary=True)
    document = _document(tmp_path / "good.docx", {"Context": ["Background"], "Decision": ["Approve rollout."]})

    assert sr.validate_slides(slides, SLIDE_PLAN) == []
    assert sr.validate_spreadsheet(spreadsheet, SHEET_PLAN) == []
    assert sr.validate_document(document, DOC_PLAN) == []


def test_validators_report_known_bad_artifacts(tmp_path: Path) -> None:
    slides = _slides(tmp_path / "bad.pptx", ["Context"])
    spreadsheet = _spreadsheet(tmp_path / "bad.xlsx", include_summary=False)
    document = _document(tmp_path / "bad.docx", {"Context": ["Background"], "Decision": []})

    assert {issue["section"] for issue in sr.validate_slides(slides, SLIDE_PLAN)} == {"Decision"}
    assert {issue["sheet"] for issue in sr.validate_spreadsheet(spreadsheet, SHEET_PLAN)} == {"Summary"}
    assert sr.validate_document(document, DOC_PLAN) == [{"type": "empty_section", "section": "Decision"}]


def test_fixers_resolve_missing_parts(tmp_path: Path) -> None:
    slides = _slides(tmp_path / "fix.pptx", ["Context"])
    slide_issues = sr.validate_slides(slides, SLIDE_PLAN)
    sr.fix_slides(slides, slide_issues, SLIDE_PLAN)
    assert sr.validate_slides(slides, SLIDE_PLAN) == []

    spreadsheet = _spreadsheet(tmp_path / "fix.xlsx", include_summary=False)
    sheet_issues = sr.validate_spreadsheet(spreadsheet, SHEET_PLAN)
    sr.fix_spreadsheet(spreadsheet, sheet_issues, SHEET_PLAN)
    assert sr.validate_spreadsheet(spreadsheet, SHEET_PLAN) == []

    document = _document(tmp_path / "fix.docx", {"Context": ["Background"], "Decision": []})
    doc_issues = sr.validate_document(document, DOC_PLAN)
    sr.fix_document(document, doc_issues, DOC_PLAN)
    assert sr.validate_document(document, DOC_PLAN) == []


def test_self_review_stops_at_iteration_cap(tmp_path: Path, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    slides = _slides(tmp_path / "cap.pptx", ["Context"])

    def no_op_fix(path, issues, plan):  # type: ignore[no-untyped-def]
        return None

    monkeypatch.setattr(sr, "fix_slides", no_op_fix)
    result = sr.self_review(slides, SLIDE_PLAN, "slides", max_iterations=2)

    assert result == {"iterations_used": 2, "issues_found": 1, "issues_resolved": 0}


def _slides(path: Path, section_titles: list[str]) -> Path:
    presentation = Presentation()
    for title in section_titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
        box.text_frame.text = title
    presentation.save(path)
    return path


def _spreadsheet(path: Path, include_summary: bool) -> Path:
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Metric", "Value"])
    data.append(["Revenue", 100])
    if include_summary:
        summary = workbook.create_sheet("Summary")
        summary.append(["Question", "Answer"])
        summary.append(["Purpose", "Review"])
    workbook.save(path)
    return path


def _document(path: Path, sections: dict[str, list[str]]) -> Path:
    document = Document()
    document.add_heading("Review", 0)
    for heading, paragraphs in sections.items():
        document.add_heading(heading, 1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.save(path)
    return path
