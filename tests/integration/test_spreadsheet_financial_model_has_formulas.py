from __future__ import annotations

import json
from pathlib import Path

from codeact.skills_loader import invoke_skill
from openpyxl import load_workbook


def test_spreadsheet_financial_model_has_formulas(tmp_path: Path) -> None:
    result = invoke_skill(
        "spreadsheet",
        {"output_dir": str(tmp_path), "mode": "financial_model", "assumptions": {"starting_revenue": 750000}},
    )
    assert result.exit_code == 0, result.stderr
    workbook_path = Path(json.loads(result.stdout)["path"])
    workbook = load_workbook(workbook_path, data_only=False)
    assert {"Assumptions", "P&L", "Balance Sheet", "Cash Flow"}.issubset(set(workbook.sheetnames))
    formulas = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert formulas
    assert "=Assumptions!B1" in formulas
