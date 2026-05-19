from __future__ import annotations

import json
from pathlib import Path

from codeact.skills_loader import invoke_skill
from openpyxl import load_workbook


def test_spreadsheet_skill_round_trip(tmp_path: Path) -> None:
    result = invoke_skill(
        "spreadsheet",
        {
            "output_dir": str(tmp_path),
            "mode": "table",
            "schema": [{"name": "Metric"}, {"name": "Value", "type": "number"}],
            "rows": [{"Metric": "Revenue", "Value": 100}, {"Metric": "Users", "Value": 42}],
            "charts": [{"type": "bar", "x": "Metric", "y": "Value"}],
        },
    )
    assert result.exit_code == 0, result.stderr
    payload = json.loads(result.stdout)
    workbook_path = Path(payload["path"])
    assert workbook_path.exists()
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == ["Data", "Charts"]
    assert workbook["Data"]["A1"].value == "Metric"
